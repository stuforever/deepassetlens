# -*- coding: utf-8 -*-
"""文档生成工具 4 件 impl（私塾切换 v3.0 R0-③ / G6，design §2.2+§三 G6）。

vendor skills/builtin/{docx,pptx,xlsx,pdf}/SKILL.md 的平台化重写：生成库
（python-docx/python-pptx/openpyxl/reportlab）留用，注册面/确认流（EXEC 两段臂）
全走平台 MCP。产物统一落 backend/data/exports/documents/，写后回读校验
（对齐 vendor SKILL.md "Verify before returning" 纪律）。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List

DOC_ROOT = Path(__file__).resolve().parents[2] / "data" / "exports" / "documents"

_NAME_RE = re.compile(r"^[\w一-鿿.\- ]{1,120}$")

_BLOCK_TYPES = {"heading", "para", "bullet", "number", "table"}


def _resolve_output(name: str, ext: str) -> Path:
    """输出名清洗：仅文件名（禁路径穿越），强制后缀，重名顺延 _2/_3。"""
    if not name or not _NAME_RE.match(name.strip()):
        raise ValueError(f"output_name 非法（仅限中英文/数字/空格/._- ，收到 {name!r}）")
    safe = Path(name.strip()).name
    if not safe.lower().endswith(ext):
        safe += ext
    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    out = DOC_ROOT / safe
    stem, n = out.stem, 2
    while out.exists():
        out = DOC_ROOT / f"{stem}_{n}{ext}"
        n += 1
    return out


def _norm_blocks(blocks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("blocks 必须为非空列表")
    for b in blocks:
        if not isinstance(b, dict) or b.get("type") not in _BLOCK_TYPES:
            raise ValueError(f"block.type 需为 {_BLOCK_TYPES} 之一，收到 {b}")
        if b["type"] == "table":
            if not isinstance(b.get("headers"), list) or not b["headers"]:
                raise ValueError("table block 需非空 headers 列表")
    return blocks


def generate_docx(blocks: List[Dict[str, Any]], output_name: str = "文档.docx") -> dict:
    """python-docx 生成 Word 文档（标题/段落/列表/表格）。"""
    from docx import Document

    blocks = _norm_blocks(blocks)
    doc = Document()
    for b in blocks:
        t = b["type"]
        if t == "heading":
            doc.add_heading(str(b.get("text", "")), level=int(b.get("level", 1) or 1))
        elif t == "para":
            doc.add_paragraph(str(b.get("text", "")))
        elif t == "bullet":
            doc.add_paragraph(str(b.get("text", "")), style="List Bullet")
        elif t == "number":
            doc.add_paragraph(str(b.get("text", "")), style="List Number")
        elif t == "table":
            tbl = doc.add_table(rows=1, cols=len(b["headers"]))
            tbl.style = "Table Grid"
            for i, h in enumerate(b["headers"]):
                tbl.rows[0].cells[i].text = str(h)
            for row in b.get("rows") or []:
                cells = tbl.add_row().cells
                for i, v in enumerate(row[:len(b["headers"])]):
                    cells[i].text = str(v)
    out = _resolve_output(output_name, ".docx")
    doc.save(str(out))
    from docx import Document as _D
    _D(str(out))  # 回读校验（损坏文件在此暴露）
    return {"file": str(out), "format": "docx"}


def generate_pptx(slides: List[Dict[str, Any]], output_name: str = "演示.pptx") -> dict:
    """python-pptx 生成演示文稿（标题页 + 内容页 bullets）。"""
    from pptx import Presentation

    if not isinstance(slides, list) or not slides:
        raise ValueError("slides 必须为非空列表")
    prs = Presentation()
    for i, s in enumerate(slides):
        if not isinstance(s, dict):
            raise ValueError(f"slide[{i}] 需为对象（title/bullets）")
        layout = prs.slide_layouts[0 if i == 0 else 1]
        slide = prs.slides.add_slide(layout)
        title = str(s.get("title", ""))
        slide.shapes.title.text = title
        bullets = [str(x) for x in (s.get("bullets") or [])]
        if bullets:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = bullets[0]
            for extra in bullets[1:]:
                p = tf.add_paragraph()
                p.text = extra
    out = _resolve_output(output_name, ".pptx")
    prs.save(str(out))
    from pptx import Presentation as _P
    _P(str(out))
    return {"file": str(out), "format": "pptx"}


def generate_xlsx(sheets: List[Dict[str, Any]], output_name: str = "表格.xlsx") -> dict:
    """openpyxl 生成工作簿（每 sheet：headers 表头 + rows 数据行）。"""
    from openpyxl import Workbook, load_workbook

    if not isinstance(sheets, list) or not sheets:
        raise ValueError("sheets 必须为非空列表")
    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets):
        if not isinstance(sh, dict) or not sh.get("headers"):
            raise ValueError(f"sheets[{i}] 需含 headers")
        ws = wb.create_sheet(title=str(sh.get("name") or f"Sheet{i + 1}")[:31])
        ws.append([str(h) for h in sh["headers"]])
        for row in sh.get("rows") or []:
            ws.append(list(row))
    out = _resolve_output(output_name, ".xlsx")
    wb.save(str(out))
    load_workbook(str(out))
    return {"file": str(out), "format": "xlsx"}


def _pdf_blocks(blocks: List[Dict[str, Any]], out: Path) -> None:
    """reportlab platypus 渲染（中文走内置 UnicodeCIDFont，免字体文件）。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle, Spacer

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = {
        "heading": ParagraphStyle("h", fontName="STSong-Light", fontSize=16,
                                  leading=22, spaceBefore=12, spaceAfter=6),
        "para": ParagraphStyle("p", fontName="STSong-Light", fontSize=11,
                               leading=17, spaceAfter=5),
        "item": ParagraphStyle("i", fontName="STSong-Light", fontSize=11,
                               leading=17, leftIndent=14, spaceAfter=3),
    }
    story = []
    num = 0
    for b in blocks:
        t = b["type"]
        if t == "heading":
            num = 0
            story.append(Paragraph(str(b.get("text", "")), styles["heading"]))
        elif t in ("para", "bullet", "number"):
            text = str(b.get("text", ""))
            if t == "bullet":
                text = f"• {text}"
            elif t == "number":
                num += 1
                text = f"{num}. {text}"
            story.append(Paragraph(text, styles["item" if t != "para" else "para"]))
        elif t == "table":
            data = [[str(h) for h in b["headers"]]]
            data += [[str(c) for c in row] for row in b.get("rows") or []]
            tb = Table(data, colWidths=[(A4[0] - 4 * cm) / max(len(b["headers"]), 1)] * len(b["headers"]))
            tb.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, "grey"),
                ("BACKGROUND", (0, 0), (-1, 0), "#eee"),
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ]))
            story.append(tb)
            story.append(Spacer(1, 8))
    doc = SimpleDocTemplate(str(out), pagesize=A4)
    doc.build(story)


def generate_pdf(blocks: List[Dict[str, Any]], output_name: str = "文档.pdf") -> dict:
    """reportlab 生成 PDF（标题/段落/列表/表格，STSong-Light 中文字体）。"""
    blocks = _norm_blocks(blocks)
    out = _resolve_output(output_name, ".pdf")
    _pdf_blocks(blocks, out)
    with open(out, "rb") as f:
        head = f.read(5)
    if head != b"%PDF-":
        raise RuntimeError("PDF 写出异常（魔数缺失）")
    return {"file": str(out), "format": "pdf"}
