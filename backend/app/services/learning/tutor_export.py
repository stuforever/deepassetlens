# -*- coding: utf-8 -*-
"""⑤b：错题本导出（engine 臂）——落文件返回 result_ref（tab_export 纪律：大对象不进对话）。"""
from __future__ import annotations

import json
import uuid
from pathlib import Path

EXPORT_ROOT = Path(__file__).resolve().parents[3] / "data" / "exports" / "wrong_book"


def export_wrong_book_file(user_id: str, format: str = "md") -> dict:
    """导出错题清单→文件→result_ref（query_result_store 同款指针语义）。"""
    from app.services.learning.learning_dao import wrong_question_export_rows
    rows = wrong_question_export_rows(user_id)
    if format not in ("md", "json", "docx"):
        return {"error": f"format 白名单 md|json|docx（收到 {format}）"}
    EXPORT_ROOT.mkdir(parents=True, exist_ok=True)
    # R5批⑰（清单安全）：user_id 来自工具调用参数，未清洗即拼导出文件名（'/' '..'
    # 空字节可穿越 EXPORT_ROOT）。折叠非安全字符为 '_'。
    import re as _re

    _uid = _re.sub(r"[^A-Za-z0-9_-]", "_", str(user_id or ""))[:64] or "anonymous"
    stem = f"wrongbook_{_uid}_{uuid.uuid4().hex[:8]}"
    if format == "json":
        path = EXPORT_ROOT / f"{stem}.json"
        path.write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str),
                        encoding="utf-8")
    elif format == "docx":
        try:
            import docx
            d = docx.Document()
            d.add_heading(f"错题本 · {user_id}", level=1)
            for r in rows:
                d.add_paragraph(str(r.get("variant_text", "")))
                d.add_paragraph(f"  状态: {r.get('status')} / 记录于 {r.get('wrong_at')}")
            path = EXPORT_ROOT / f"{stem}.docx"
            d.save(str(path))
        except Exception as e:
            return {"error": f"docx 导出失败: {e}"}
    else:
        lines = [f"# 错题本 · {user_id}", ""]
        for r in rows:
            lines.append(f"- [{r.get('status')}] {r.get('variant_text', '')}")
            if r.get("error_context"):
                lines.append(f"  - 上下文: {r['error_context']}")
        path = EXPORT_ROOT / f"{stem}.md"
        path.write_text("\n".join(lines), encoding="utf-8")

    # result_ref 沉降（⑤b 铁律③）：内容指针+摘要视图
    from app.services import query_result_store as _qrs
    ref = _qrs.put({"rows": [{"file": str(path), "count": len(rows)}],
                    "row_count": 1, "columns": ["file", "count"]})
    return {"result_ref": ref, "file": str(path), "count": len(rows), "format": format}
