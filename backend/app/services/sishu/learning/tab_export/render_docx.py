# -*- coding: utf-8 -*-
# [sishu port] v4批6 机械复制自 vendor deeptutor/learning/tab_export/render_docx.py（1:1 语义，仅 import 改写）。
# 批16 vendor 物理删除后的存活拷贝；上游修订需回灌本件（复刻纪律）。
"""DOCX 渲染器（python-docx，T15）。

中文字体走 eastAsia 设置（``qn("w:eastAsia")``，先例 mother_question.py
L1288 同款 oxml 操作）：西文 Calibri + 中文宋体，Word/WPS 通用。
"""

from __future__ import annotations

import io
from typing import Any

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from app.services.sishu.learning.tab_export.model import ExportDoc, Section, format_answer

__all__ = ["render_docx"]

_EASTASIA_FONT = "SimSun"
_LATIN_FONT = "Calibri"

# XML 1.0 合法字符（document.xml 序列化拒绝控制符 → ValueError/500，
# T21 e2e 实测七上数学 OCR 原文含 \x0b）——渲染前剥离
_XML_ILLEGAL = re.compile(
    "[^\t\n\r\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]"
)


def _xml_safe(text: str) -> str:
    """剥除 XML 1.0 非法字符（OCR 文本常见 \x0b/\x0c/\x00 等控制符）。"""
    return _XML_ILLEGAL.sub("", text)


def _style_run(run) -> None:
    """每个 run 统一设置西文 + eastAsia 中文字体（含表格/列表内文本）。"""
    run.font.name = _LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _EASTASIA_FONT)


def _add_text(doc: Document, text: str, *, size: float = 10.5,
              bold: bool = False, align=None, color: str | None = None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(_xml_safe(str(text)))
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    _style_run(run)
    return p


def _add_table(doc: Document, table: dict[str, Any]) -> None:
    headers = [_xml_safe(str(h)) for h in table.get("headers") or []]
    rows = [[_xml_safe(str(c)) for c in (r or [])] for r in table.get("rows") or []]
    if not headers and not rows:
        return
    n_cols = len(headers) or max((len(r) for r in rows), default=0)
    if not n_cols:
        return
    flow_table = doc.add_table(rows=0, cols=n_cols)
    flow_table.style = "Table Grid"
    if headers:
        cells = flow_table.add_row().cells
        for j, h in enumerate(headers[:n_cols]):
            cells[j].text = ""
            _style_run(cells[j].paragraphs[0].add_run(h))
            for run in cells[j].paragraphs[0].runs:
                run.bold = True
    for row in rows:
        cells = flow_table.add_row().cells
        for j in range(n_cols):
            cells[j].text = ""
            _style_run(cells[j].paragraphs[0].add_run(row[j] if j < len(row) else ""))


def _add_section(doc: Document, section: Section) -> None:
    if section.heading:
        _add_text(doc, section.heading, size=12.5, bold=True)
    for text in section.paragraphs:
        _add_text(doc, str(text))
    for item in section.list_items or []:
        p = _add_text(doc, f"• {item}")
        p.paragraph_format.left_indent = Pt(12)
    if section.table:
        _add_table(doc, section.table)
    for i, q in enumerate(section.questions or [], start=1):
        _add_text(doc, f"{i}. {q.get('ask', '')}", bold=False)
        options = q.get("options") or []
        for j, opt in enumerate(options):
            p = _add_text(doc, f"{'ABCDEFGH'[j] if j < 8 else j}. {opt}")
            p.paragraph_format.left_indent = Pt(12)
        answer = q.get("answer")
        # 审查 R1-建议1：填空/简答题（无 options）也渲染答案（家长版判例 3）
        answer_label = format_answer(answer)
        if answer_label is not None:
            _add_text(doc, f"答案：{answer_label}", color="2e7d32")
        why = q.get("why")
        if why:
            _add_text(doc, f"解析：{why}", size=9.5, color="666666")


def render_docx(doc: ExportDoc) -> bytes:
    """ExportDoc → DOCX 字节流（zip → b"PK" 文件头）。"""
    document = Document()
    if doc.title:
        _add_text(document, doc.title, size=16, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for m in doc.meta or []:
        _add_text(document, f"{m.get('label', '')}：{m.get('value', '')}",
                  size=9, color="555555")
    for section in doc.sections or []:
        _add_section(document, section)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
