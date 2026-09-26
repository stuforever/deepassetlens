# -*- coding: utf-8 -*-
"""文档生成 MCP 4 件测试（切换 R0-③ / G6）：生成正确性 + 回读校验 + 两段臂确认流。"""
from __future__ import annotations

import os
from pathlib import Path

import pytest

from app.services import doc_tools

_BLOCKS = [
    {"type": "heading", "text": "季度报告", "level": 1},
    {"type": "para", "text": "本季度营收增长显著。"},
    {"type": "bullet", "text": "要点一"},
    {"type": "number", "text": "步骤一"},
    {"type": "number", "text": "步骤二"},
    {"type": "table", "headers": ["指标", "值"], "rows": [["营收", "1.2M"]]},
]


@pytest.fixture(autouse=True)
def _tmp_doc_root(tmp_path, monkeypatch):
    monkeypatch.setattr(doc_tools, "DOC_ROOT", tmp_path / "documents")
    return tmp_path


def test_generate_docx():
    out = doc_tools.generate_docx(_BLOCKS, "测试文档")
    p = Path(out["file"])
    assert p.exists() and p.suffix == ".docx"
    from docx import Document
    d = Document(str(p))
    assert any("季度报告" in para.text for para in d.paragraphs)
    assert len(d.tables) == 1


def test_generate_pptx():
    out = doc_tools.generate_pptx([
        {"title": "封面", "bullets": []},
        {"title": "正文页", "bullets": ["要点A", "要点B"]},
    ], "测试演示")
    p = Path(out["file"])
    assert p.exists() and p.suffix == ".pptx"
    from pptx import Presentation
    prs = Presentation(str(p))
    assert len(prs.slides) == 2
    assert prs.slides[1].shapes.title.text == "正文页"


def test_generate_xlsx():
    out = doc_tools.generate_xlsx([
        {"name": "成绩", "headers": ["姓名", "分数"], "rows": [["张三", 95]]},
    ], "测试表格")
    p = Path(out["file"])
    assert p.exists() and p.suffix == ".xlsx"
    from openpyxl import load_workbook
    wb = load_workbook(str(p))
    ws = wb["成绩"]
    assert ws.cell(row=1, column=1).value == "姓名"
    assert ws.cell(row=2, column=2).value == 95


def test_generate_pdf():
    out = doc_tools.generate_pdf(_BLOCKS, "测试PDF")
    p = Path(out["file"])
    assert p.exists() and p.suffix == ".pdf"
    assert p.read_bytes()[:5] == b"%PDF-"


def test_name_sanitized_and_ext_enforced():
    out = doc_tools.generate_docx([{"type": "para", "text": "x"}], "无后缀名")
    assert out["file"].endswith("无后缀名.docx")
    with pytest.raises(ValueError):
        doc_tools.generate_docx([{"type": "para", "text": "x"}], "../evil.docx")
    with pytest.raises(ValueError):
        doc_tools.generate_docx([{"type": "para", "text": "x"}], "a/b.docx")
    with pytest.raises(ValueError):
        doc_tools.generate_docx([], "x.docx")  # 空 blocks
    with pytest.raises(ValueError):
        doc_tools.generate_docx([{"type": "unknown", "text": "x"}], "x.docx")
    # 重名顺延 _2
    doc_tools.generate_xlsx([{"name": "s", "headers": ["h"], "rows": []}], "重名表")
    out2 = doc_tools.generate_xlsx([{"name": "s", "headers": ["h"], "rows": []}], "重名表")
    assert Path(out2["file"]).stem.endswith("_2")


def test_mcp_doc_tools_two_arm():
    """两段臂：无 token=pending；错 token=denied；对 token=执行。"""
    import asyncio
    from app.mcp_server import mcp

    tools = {t.name: t for t in asyncio.run(mcp.list_tools())}
    assert {"generate_docx", "generate_pptx", "generate_xlsx", "generate_pdf"} <= set(tools)

    from app.services import memory_runtime as mr
    from app.services import tool_confirm
    token = mr.set_runtime("sishu", "tester", "sess-t")
    blocks = [{"type": "para", "text": "两段臂"}]
    bad = tool_confirm.issue("generate_docx", {"blocks": blocks, "output_name": "别的"}, "tester")

    import app.mcp_server as ms
    try:
        # 无 token → 预检臂
        r1 = ms.generate_docx(blocks=blocks, output_name="两段臂")
        assert r1["status"] == "pending_confirmation" and r1["confirm_token"]
        # 错 token → 拒
        r2 = ms.generate_docx(blocks=blocks, output_name="两段臂", confirm_token=bad)
        assert r2["status"] == "denied"
        # 对 token → 执行（产物落临时 DOC_ROOT）
        good = r1["confirm_token"]
        r3 = ms.generate_docx(blocks=blocks, output_name="两段臂", confirm_token=good)
        assert Path(r3["file"]).exists()
    finally:
        mr.set_runtime_token(token)
